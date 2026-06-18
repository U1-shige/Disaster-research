"""
DocuWorks 変換 CLI
Power Automate が OneDrive に保存した突合結果 Excel を読み込み、
「真」「要確認」の案件を DocuWorks (.xdw) に変換する。

使い方:
  python main_xdw.py \
    --excel   突合結果.xlsx \
    --pdf-dir C:\\Users\\you\\Downloads\\二次調査 \
    --pages   3,5 \
    --output  xdw_out
"""
import argparse
import os
import sys
from pathlib import Path

import openpyxl

from xdw_converter import DEFAULT_XDWAPI_DLL, convert_to_xdw

TARGET_JUDGMENTS = {"真", "要確認"}


def find_pdf(pdf_dir: str, folder_num: str, filename: str) -> str | None:
    """フォルダ番号とファイル名から PDF を探す。"""
    candidate = Path(pdf_dir) / str(folder_num) / filename
    if candidate.exists():
        return str(candidate)
    # ファイル名だけで再検索（パス区切り違い対策）
    for p in Path(pdf_dir).rglob(filename):
        return str(p)
    return None


def run(excel_path: str, pdf_dir: str, pages: list[int], output_dir: str, xdwapi_dll: str):
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col = {name: idx for idx, name in enumerate(headers)}

    required = {"フォルダ番号", "証拠ファイル名", "判定"}
    missing = required - set(col.keys())
    if missing:
        print(f"[エラー] Excel に必要な列がありません: {missing}", file=sys.stderr)
        sys.exit(1)

    results = {"成功": 0, "スキップ": 0, "エラー": 0}

    # DocuWorks変換列を探す（なければ追加）
    xdw_col_name = "DocuWorks変換"
    if xdw_col_name not in col:
        ws.cell(row=1, column=len(headers) + 1, value=xdw_col_name)
        col[xdw_col_name] = len(headers)
        headers.append(xdw_col_name)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        judgment = row[col["判定"]].value
        if judgment not in TARGET_JUDGMENTS:
            results["スキップ"] += 1
            continue

        folder_num = row[col["フォルダ番号"]].value
        evidence_file = row[col["証拠ファイル名"]].value

        pdf_path = find_pdf(pdf_dir, folder_num, evidence_file)
        if pdf_path is None:
            print(f"[警告] PDF が見つかりません: フォルダ {folder_num} / {evidence_file}")
            ws.cell(row=row_idx, column=col[xdw_col_name] + 1, value="ファイルなし")
            results["エラー"] += 1
            continue

        case_output_dir = str(Path(output_dir) / str(folder_num))
        Path(case_output_dir).mkdir(parents=True, exist_ok=True)

        try:
            xdw_path = convert_to_xdw(pdf_path, pages, case_output_dir, xdwapi_dll)
            ws.cell(row=row_idx, column=col[xdw_col_name] + 1, value="○")
            print(f"[完了] {xdw_path}")
            results["成功"] += 1
        except Exception as e:
            print(f"[エラー] フォルダ {folder_num} / {evidence_file}: {e}", file=sys.stderr)
            ws.cell(row=row_idx, column=col[xdw_col_name] + 1, value=f"エラー: {e}")
            results["エラー"] += 1

    wb.save(excel_path)
    print(
        f"\n完了: 成功 {results['成功']} 件 / スキップ {results['スキップ']} 件 / エラー {results['エラー']} 件"
    )


def main():
    parser = argparse.ArgumentParser(description="突合結果Excelを元にDocuWorks変換を実行する")
    parser.add_argument("--excel", required=True, help="突合結果 Excel ファイルのパス")
    parser.add_argument("--pdf-dir", required=True, help="二次調査フォルダのローカルパス（番号フォルダを含む）")
    parser.add_argument(
        "--pages",
        required=True,
        help="出力するページ番号（カンマ区切り、例: 3,5）",
    )
    parser.add_argument("--output", required=True, help="DocuWorks 出力フォルダ")
    parser.add_argument(
        "--xdwapi-dll",
        default=DEFAULT_XDWAPI_DLL,
        help=f"XDWAPI.dll のパス（省略時: {DEFAULT_XDWAPI_DLL}）",
    )
    args = parser.parse_args()

    pages = [int(p.strip()) for p in args.pages.split(",")]

    run(
        excel_path=args.excel,
        pdf_dir=args.pdf_dir,
        pages=pages,
        output_dir=args.output,
        xdwapi_dll=args.xdwapi_dll,
    )


if __name__ == "__main__":
    main()
